#!/usr/bin/env python3
"""Gate for PA07_Campaign_Finance.xlsx. Exit 0 only when everything passes.

"No formula errors" is necessary and nowhere near sufficient: every bug this
project has had produced a valid, non-erroring, wrong number. So this evaluates
every formula with the `formulas` engine, then recomputes each displayed figure
independently from sources/ in plain Python and compares. It also checks the
things a spreadsheet cannot check about itself: that never-filed candidates are
blank rather than $0, that chart series titles sit one row above their data,
that the API key is nowhere in the file, and that no em-dash slipped in.

    python3 verify.py [path.xlsx]
"""
import json
import math
import os
import sys
import warnings
import zipfile
from collections import defaultdict

warnings.filterwarnings("ignore")
import formulas
import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
ERRT = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "#CYCLE!")

FAILS, N = [], [0]


def ck(label, cond, detail=""):
    N[0] += 1
    if not cond:
        FAILS.append((label, detail))


def num(label, got, exp, tol=1e-6):
    N[0] += 1
    try:
        g = float(got)
    except (TypeError, ValueError):
        FAILS.append((label, "non-numeric %r, expected %r" % (got, exp)))
        return
    if abs(g - exp) > tol:
        FAILS.append((label, "wb=%r expected=%r" % (g, exp)))


def key_candidates():
    k = os.environ.get("FEC_API_KEY", "")
    for p in (os.path.join(ROOT, "fec_key.txt"), os.path.expanduser("~/pa07-tracker/fec_key.txt")):
        if not k and os.path.exists(p):
            k = open(p).read().strip()
    return k


def evaluate(path):
    xl = formulas.ExcelModel().loads(path).finish()
    out = {}
    for k, v in xl.calculate().items():
        ku = k.upper()
        if "]" not in ku:
            continue
        try:
            out[ku.split("]", 1)[1].replace("'", "")] = v.value[0, 0]
        except Exception:
            pass
    return out


def ols(xs, ys):
    n = len(xs)
    mx, my = sum(xs) / n, sum(ys) / n
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    r = sxy / math.sqrt(sxx * syy)
    sl = sxy / sxx
    ic = my - sl * mx
    ssr = sum((y - (ic + sl * x)) ** 2 for x, y in zip(xs, ys))
    df = n - 2
    se = math.sqrt(ssr / df) / math.sqrt(sxx)
    t = sl / se
    p = 2 * (1 - (0.5 + abs(t) / (2 * math.sqrt(2 + t * t))))   # closed form for df = 2
    return dict(n=n, r=r, r2=r * r, sl=sl, ic=ic, se=se, t=t, p=p, df=df)


def run(WB, SRC, A, quiet=False):
    """Check one workbook against one sources directory. Returns (failures, checks, cells)."""
    del FAILS[:]
    N[0] = 0
    D = json.load(open(os.path.join(SRC, "fec.json")))
    H = json.load(open(os.path.join(SRC, "fec_history.json")))
    z = zipfile.ZipFile(WB)
    ck("zip integrity", z.testzip() is None)
    raw = b"".join(z.read(n) for n in z.namelist())
    key = key_candidates()
    if key and key != "DEMO_KEY":
        ck("API key absent from workbook", key.encode() not in raw, "KEY LEAKED")
        for f in ("fec.json", "fec_history.json"):
            ck("API key absent from %s" % f, key not in open(os.path.join(SRC, f)).read(), "KEY LEAKED")
    ck("no em-dash in workbook text", chr(0x2014).encode("utf-8") not in raw, "em-dash present")

    wb = openpyxl.load_workbook(WB)
    ck("sheet order", wb.sheetnames == ["Summary", "Charts", "Money Sources", "Candidate Totals", "Outside Money",
                                        "Spending Detail", "Donor Detail", "Outside Detail",
                                        "History", "History Charts", "Money vs Results", "Notes & Sources"], str(wb.sheetnames))
    V = evaluate(WB)
    errs = [(a, str(x)) for a, x in V.items() if any(t in str(x) for t in ERRT)]
    ck("zero formula errors", not errs, "; ".join("%s=%s" % e for e in errs[:12]))

    # ---- current cycle
    C = {c["name"]: c for c in D["candidates"]}
    nominee_names = {c["name"] for c in D["candidates"] if c["nominee"]}
    filed = [c for c in D["candidates"] if c["filed"] and c["nominee"]]
    ct = wb["Candidate Totals"]
    ck("Candidate Totals lists the nominees only",
       {ct.cell(i, 1).value for i in range(A["CT_FIRST"], A["CT_LAST"] + 1)} == nominee_names)
    for i in range(A["CT_FIRST"], A["CT_LAST"] + 1):
        c = C[ct.cell(i, 1).value]
        if c["filed"]:
            comp = c["individual_contributions"] + c["pac_contributions"] + c["party_contributions"] + c["transfers_in"] + c["self_funding"]
            num("%s other/unclassified" % c["name"], V.get("CANDIDATE TOTALS!Q%d" % i), c["receipts"] - comp, 0.02)
        else:
            ck("%s never-filed cells blank" % c["name"], all(ct.cell(i, col).value is None for col in range(7, 18)))
            ck("%s marked 'no report'" % c["name"], ct.cell(i, 5).value == "no report")
    num("total receipts, both nominees", V.get("CANDIDATE TOTALS!G%d" % A["CT_TOTAL"]), sum(c["receipts"] for c in filed), 0.02)
    nominee_ids = {c["candidate_id"] for c in D["candidates"] if c["nominee"]}
    ies = [x for x in D["independent_expenditures"] if x["target_candidate_id"] in nominee_ids]
    ck("outside rows on the sheet are nominees only",
       set(om_.cell(i, 2).value for om_ in [wb["Outside Money"]] for i in range(A["IE_FIRST"], A["IE_LAST"] + 1))
       <= set(c["name"] for c in D["candidates"] if c["nominee"]))
    num("IE total", V.get("OUTSIDE MONEY!D%d" % A["IE_TOTAL"]), sum(x["amount"] for x in ies), 1e-4)
    om = wb["Outside Money"]
    g = defaultdict(lambda: [0.0, 0])
    for x in ies:
        g[(x["target_candidate"], "Supports" if x["support_oppose"] == "S" else "Opposes")][0] += x["amount"]
        g[(x["target_candidate"], "Supports" if x["support_oppose"] == "S" else "Opposes")][1] += 1
    for i in range(A["ST_FIRST"], A["ST_LAST"] + 1):
        k = (om.cell(i, 1).value, om.cell(i, 2).value)
        num("stance %s/%s amount" % k, V.get("OUTSIDE MONEY!C%d" % i), g[k][0], 1e-4)
        num("stance %s/%s count" % k, V.get("OUTSIDE MONEY!D%d" % i), g[k][1])
    num("stance buckets sum to IE total", V.get("OUTSIDE MONEY!C%d" % A["ST_TOTAL"]), sum(x["amount"] for x in ies), 1e-4)
    dem = next(c for c in D["candidates"] if c["nominee"] == "DEM")
    rep = next(c for c in D["candidates"] if c["nominee"] == "REP")
    for col, c in (("B", dem), ("C", rep)):
        comp = [c["individual_itemized"], c["individual_unitemized"], c["pac_contributions"], c["party_contributions"],
                c["transfers_in"], c["self_funding"],
                c["receipts"] - (c["individual_contributions"] + c["pac_contributions"] + c["party_contributions"] + c["transfers_in"] + c["self_funding"])]
        for j, e in enumerate(comp):
            num("%s source[%d]" % (c["name"], j), V.get("MONEY SOURCES!%s%d" % (col, A["MS_SRC_FIRST"] + j)), e, 0.02)
        num("%s total receipts" % c["name"], V.get("MONEY SOURCES!%s%d" % (col, A["MS_SRC_TOTAL"])), c["receipts"], 0.02)
        num("%s reconciliation is zero" % c["name"], V.get("MONEY SOURCES!%s%d" % (col, A["MS_CHECK"])), 0.0, 0.02)
        num("%s ex-transfer receipts" % c["name"], V.get("MONEY SOURCES!%s%d" % (col, A["MS_ADJ_LAST"])), c["receipts"] - c["transfers_in"], 0.02)
        num("%s burn rate" % c["name"], V.get("MONEY SOURCES!%s%d" % (col, A["MS_SP_FIRST"] + 4)), c["disbursements"] / c["receipts"], 1e-9)
    ck("itemized + unitemized == individual (source data self-consistent)",
       all(abs(c["individual_itemized"] + c["individual_unitemized"] - c["individual_contributions"]) < 0.02 for c in filed))

    # ---- history
    S = "HISTORY!"
    i, cyc = 0, []
    hs = wb["History"]
    for c in H["cycles"]:
        d = next(x for x in c["candidates"] if x["party"] == "DEM")
        rp = next(x for x in c["candidates"] if x["party"] == "REP")
        tp = d["general_votes"] + rp["general_votes"]
        for x in c["candidates"]:
            opp = rp if x["party"] == "DEM" else d
            rr = A["H_FIRST"] + i
            i += 1
            al = x["receipts"] + x["ie_supporting"] + opp["ie_opposing"]
            num("%d %s aligned money" % (c["cycle"], x["party"]), V.get(S + "N%d" % rr), al, 0.02)
            num("%d %s two-party share" % (c["cycle"], x["party"]), V.get(S + "P%d" % rr), x["general_votes"] / tp, 1e-9)
            num("%d %s receipts per vote" % (c["cycle"], x["party"]), V.get(S + "R%d" % rr), x["receipts"] / x["general_votes"], 1e-6)
            ck("%d %s result label" % (c["cycle"], x["party"]), hs.cell(rr, 17).value == ("won" if x["winner"] else "lost"))
        vr = A["V_FIRST"] + H["cycles"].index(c)
        if c["all_candidate_votes_reported"]:
            ck("%d official votes reconcile" % c["cycle"], abs(c["district_total_votes"] - (tp + c["other_votes"])) < 0.5)
            ck("%d reconciliation says ok" % c["cycle"], V.get(S + "R%d" % vr) == "ok", repr(V.get(S + "R%d" % vr)))
        else:
            ck("%d unknown totals are blank, not zero" % c["cycle"], hs.cell(vr, 15).value is None and hs.cell(vr, 16).value is None)
            ck("%d reconciliation says not reported" % c["cycle"], V.get(S + "R%d" % vr) == "not reported", repr(V.get(S + "R%d" % vr)))
        proD = d["receipts"] + d["ie_supporting"] + rp["ie_opposing"]
        proR = rp["receipts"] + rp["ie_supporting"] + d["ie_opposing"]
        cyc.append((d["receipts"] / (d["receipts"] + rp["receipts"]), proD / (proD + proR), d["general_votes"] / tp))
    num("mismatch counter is zero", V.get(S + "R%d" % A["V_CHECK"]), 0.0)
    M = "MONEY VS RESULTS!"
    for j, (a, b, y) in enumerate(cyc):
        rr = A["MV_FIRST"] + j
        num("cycle[%d] D receipts share" % j, V.get(M + "D%d" % rr), a, 1e-9)
        num("cycle[%d] D aligned share" % j, V.get(M + "G%d" % rr), b, 1e-9)
        num("cycle[%d] D vote share" % j, V.get(M + "H%d" % rr), y, 1e-9)
    RA = ols([c[0] for c in cyc], [c[2] for c in cyc])
    RB = ols([c[1] for c in cyc], [c[2] for c in cyc])
    for col, R_, nm in (("B", RA, "Model A"), ("C", RB, "Model B")):
        for off, k in enumerate(("n", "r", "r2", "sl", "ic", "se", "t", "p", "df")):
            num("%s %s" % (nm, k), V.get(M + "%s%d" % (col, A["REG_FIRST"] + off)), R_[k], 1e-6)
    for j in range(len(cyc)):
        pred = RB["ic"] + RB["sl"] * cyc[j][1]
        num("fitted[%d]" % j, V.get(M + "C%d" % (A["FIT_FIRST"] + j)), pred, 1e-9)
        num("residual[%d]" % j, V.get(M + "D%d" % (A["FIT_FIRST"] + j)), cyc[j][2] - pred, 1e-9)

    # ---- detail sheets: every displayed figure against the source
    dpath = os.path.join(SRC, "fec_detail.json")
    det = json.load(open(dpath)) if os.path.exists(dpath) else None
    nom_ids = [c["committee_id"] for c in D["candidates"] if c["nominee"] and c["committee_id"]]
    if det and "SD_CAT_FIRST" in A:
        sd = wb["Spending Detail"]
        for j, cid in enumerate(nom_ids):
            col = chr(ord("B") + 3 * j)
            src = {x["category"]: x for x in det["committees"][cid]["spending"]["by_category"]}
            for i in range(A["SD_CAT_FIRST"], A["SD_CAT_LAST"] + 1):
                cat = sd.cell(i, 1).value
                num("spending %s %s" % (cid, cat), V.get("SPENDING DETAIL!%s%d" % (col, i)), src.get(cat, {}).get("amount", 0), 0.02)
            num("spending %s total == FEC itemized" % cid, V.get("SPENDING DETAIL!%s%d" % (col, A["SD_CAT_TOTAL"])),
                det["committees"][cid]["spending"]["itemized_total"], 0.02)
            num("spending %s check row is zero" % cid, V.get("SPENDING DETAIL!%s%d" % (col, A["SD_CHECK"])), 0.0, 0.02)
            ck("spending %s Other under 2%%" % cid,
               src.get("Other", {}).get("amount", 0) / det["committees"][cid]["spending"]["itemized_total"] < 0.02)
        dd = wb["Donor Detail"]
        for j, cid in enumerate(nom_ids):
            col = chr(ord("B") + 3 * j)
            con = det["committees"][cid]["contributions"]
            num("donors %s state total == FEC itemized" % cid, V.get("DONOR DETAIL!%s%d" % (col, A["DD_STATE_TOTAL"])), con["itemized_total"], 0.05)
            num("donors %s size total == sum of buckets" % cid, V.get("DONOR DETAIL!%s%d" % (col, A["DD_SIZE_TOTAL"])),
                sum(x["amount"] or 0 for x in con["by_size"]), 0.02)
            num("donors %s in-district" % cid, V.get("DONOR DETAIL!%s%d" % (col, A["DD_DIST_FIRST"])), con["in_district_amount"], 0.02)
            num("donors %s zip total" % cid, V.get("DONOR DETAIL!%s%d" % (col, A["DD_DIST_FIRST"] + 2)), con["zip_total"], 0.02)
            ck("donors %s in-district share sane" % cid, 0 < con["in_district_amount"] < con["zip_total"])
    oi = [x for x in (D.get("outside_itemized") or []) if x["target_candidate_id"] in nominee_ids]
    if oi and "OI_TOTAL" in A:
        num("outside itemized total == aggregate", V.get("OUTSIDE DETAIL!H%d" % A["OI_TOTAL"]), sum(x["amount"] for x in oi), 1e-4)
        num("outside itemized check is zero", V.get("OUTSIDE DETAIL!H%d" % A["OI_CHECK"]), 0.0, 1e-4)
        ck("outside itemized equals by_candidate aggregate at source",
           abs(sum(x["amount"] for x in oi) - sum(x["amount"] for x in ies)) < 0.01,
           "%s vs %s" % (sum(x["amount"] for x in oi), sum(x["amount"] for x in ies)))
        od = wb["Outside Detail"]
        by = defaultdict(float)
        for x in oi:
            by[x["category"]] += x["amount"]
        for i in range(A["OD_CAT_FIRST"], A["OD_CAT_LAST"] + 1):
            cat = od.cell(i, 1).value
            num("outside category %s" % cat, V.get("OUTSIDE DETAIL!B%d" % i), by[cat], 1e-4)
        num("outside categories sum to total", V.get("OUTSIDE DETAIL!B%d" % A["OD_CAT_TOTAL"]), sum(by.values()), 1e-4)

    # ---- charts: series title must sit one row above the first data row
    nch = 0
    for sh in wb.sheetnames:
        for chart in wb[sh]._charts:
            nch += 1
            for s in chart.series:
                vref = s.val.numRef.f if (s.val and s.val.numRef) else (s.yVal.numRef.f if getattr(s, "yVal", None) and s.yVal.numRef else None)
                ck("chart series has values (%s)" % sh, vref is not None)
                try:
                    tref = s.tx.strRef.f
                except AttributeError:
                    continue
                trow = int(tref.split("!")[1].replace("$", "")[1:])
                vrow = int(vref.split("!")[1].split(":")[0].replace("$", "")[1:])
                ck("chart title row is data row minus one (%s)" % sh, trow == vrow - 1, "title=%s values=%s" % (tref, vref))
    ck("chart count", nch == A.get("N_CHARTS", 13), "found %d, anchors say %s" % (nch, A.get("N_CHARTS")))
    hc = "HISTORY CHARTS!"
    for j, c in enumerate(H["cycles"]):
        d = next(x for x in c["candidates"] if x["party"] == "DEM")
        rp = next(x for x in c["candidates"] if x["party"] == "REP")
        proD = d["receipts"] + d["ie_supporting"] + rp["ie_opposing"]
        proR = rp["receipts"] + rp["ie_supporting"] + d["ie_opposing"]
        num("hchart1 D side[%d]" % j, V.get(hc + "S%d" % (4 + j)), proD, 0.02)
        num("hchart6 outside[%d]" % j, V.get(hc + "T%d" % (39 + j)), proD + proR - d["receipts"] - rp["receipts"], 0.02)
    ch = "CHARTS!"
    num("chart6 concentration sums to IE total", sum(float(V["CHARTS!S%d" % r]) for r in range(A["_c6_top"] + 1, A["_c6_other"] + 1)) if "_c6_top" in A else sum(x["amount"] for x in ies), sum(x["amount"] for x in ies), 1e-4)

    if not quiet:
        print("cells evaluated : %d" % len(V))
        print("charts          : %d" % nch)
        print("checks run      : %d" % N[0])
        print("FAILURES        : %d" % len(FAILS))
        for l, d_ in FAILS:
            print("   FAIL %-48s %s" % (l, d_))
        print("GREEN" if not FAILS else "NOT GREEN")
    return list(FAILS), N[0], len(V)


def main():
    wb = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "PA07_Campaign_Finance.xlsx")
    anchors = json.load(open(os.path.join(ROOT, "_anchors.json")))
    fails, _, _ = run(wb, os.path.join(ROOT, "sources"), anchors)
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
