"""verify.py must pass a clean build and must catch a corrupted one."""
import json
import os

import openpyxl
import support


def test_clean_build_is_green():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        v = support.load("verify")
        fails, n, cells = v.run(out, src, A, quiet=True)
        assert not fails, fails
        assert n > 150 and cells > 1000


def test_gate_catches_a_zero_where_a_blank_belongs():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        ws = wb["Candidate Totals"]
        row = next(r for r in range(A["CT_FIRST"], A["CT_LAST"] + 1) if ws.cell(r, 5).value == "no report")
        ws.cell(row, 7).value = 0
        wb.save(out)
        v = support.load("verify")
        fails, _, _ = v.run(out, src, A, quiet=True)
        assert any("never-filed" in f[0] for f in fails), fails


def test_gate_catches_a_shortened_regression_range():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        ws = wb["Money vs Results"]
        c = ws.cell(A["REG_FIRST"] + 1, 2)
        f, l = A["MV_FIRST"], A["MV_LAST"]
        c.value = c.value.replace("D%d:D%d" % (f, l), "D%d:D%d" % (f, l - 1)).replace("H%d:H%d" % (f, l), "H%d:H%d" % (f, l - 1))
        wb.save(out)
        v = support.load("verify")
        fails, _, _ = v.run(out, src, A, quiet=True)
        assert any("Model A r" == f_[0] for f_ in fails), fails


def test_gate_catches_a_chart_title_off_by_one():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        ch = wb["Charts"]._charts[0]
        ch.series[0].val.numRef.f = "'Charts'!$S$3:$S$6"
        wb.save(out)
        v = support.load("verify")
        fails, _, _ = v.run(out, src, A, quiet=True)
        assert any("title row" in f_[0] for f_ in fails), fails
