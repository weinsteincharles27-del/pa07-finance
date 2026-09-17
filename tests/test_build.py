"""build_workbook.py against the frozen sources."""
import openpyxl
import support


def test_never_filed_candidates_are_blank_not_zero():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        fec = support.fixture("fec.json")
        ws = openpyxl.load_workbook(out)["Candidate Totals"]
        names = {ws.cell(r, 1).value: r for r in range(A["CT_FIRST"], A["CT_LAST"] + 1)}
        for c in fec["candidates"]:
            r = names[c["name"]]
            if not c["filed"]:
                assert all(ws.cell(r, col).value is None for col in range(7, 18)), c["name"]
                assert ws.cell(r, 5).value == "no report"


def test_sheet_order_and_chart_count():
    with support.sandbox() as (src, work):
        out, _ = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["Summary", "Charts", "Money Sources", "Candidate Totals", "Outside Money",
                                 "Spending Detail", "Donor Detail", "Outside Detail",
                                 "History", "History Charts", "Money vs Results", "Notes & Sources"]
        assert sum(len(wb[s]._charts) for s in wb.sheetnames) == 16


def test_chart_series_titles_sit_one_row_above_their_data():
    with support.sandbox() as (src, work):
        out, _ = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        for s in wb.sheetnames:
            for ch in wb[s]._charts:
                for ser in ch.series:
                    try:
                        tref = ser.tx.strRef.f
                    except AttributeError:
                        continue
                    vref = ser.val.numRef.f
                    trow = int(tref.split("!")[1].replace("$", "")[1:])
                    vrow = int(vref.split("!")[1].split(":")[0].replace("$", "")[1:])
                    assert trow == vrow - 1, (s, tref, vref)


def test_cash_lookups_are_keyed_by_name_not_row():
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        ws = openpyxl.load_workbook(out)["Money Sources"]
        f = ws.cell(A["MS_SP_FIRST"] + 2, 2).value       # cash on hand, Democratic nominee
        assert "MATCH(" in f and "INDEX(" in f and "Bob Brooks" in f


def test_no_em_dashes_anywhere():
    with support.sandbox() as (src, work):
        out, _ = support.built(src, work)
        assert chr(0x2014).encode("utf-8") not in open(out, "rb").read()


def test_detail_sheets_reconcile_to_source():
    """Category rows sum to FEC's itemized total; itemized outside rows sum to the aggregate."""
    import json
    with support.sandbox() as (src, work):
        out, A = support.built(src, work)
        det = json.load(open(src + "/fec_detail.json"))
        fec = json.load(open(src + "/fec.json"))
        wb = openpyxl.load_workbook(out)
        sd = wb["Spending Detail"]
        for j, cid in enumerate(c["committee_id"] for c in fec["candidates"] if c["nominee"]):
            col = 2 + 3 * j
            total = sum(sd.cell(r, col).value or 0 for r in range(A["SD_CAT_FIRST"], A["SD_CAT_LAST"] + 1))
            assert abs(total - det["committees"][cid]["spending"]["itemized_total"]) < 0.02
        od = wb["Outside Detail"]
        total = sum(od.cell(r, 8).value or 0 for r in range(A["OI_FIRST"], A["OI_LAST"] + 1))
        assert abs(total - sum(x["amount"] for x in fec["independent_expenditures"])) < 0.01


def test_workbook_builds_without_detail_file():
    """The detail is optional: a missing fec_detail.json yields a placeholder, not a crash."""
    import json, os
    with support.sandbox() as (src, work):
        os.remove(os.path.join(src, "fec_detail.json"))
        out, A = support.built(src, work)
        wb = openpyxl.load_workbook(out)
        assert "Spending Detail" in wb.sheetnames
        assert A["N_CHARTS"] == 13 and sum(len(wb[s]._charts) for s in wb.sheetnames) == 13
